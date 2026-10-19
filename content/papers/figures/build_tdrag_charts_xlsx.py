"""Build tax-drag-charts.xlsx from the two chart CSVs (export_tdrag_charts.py), mirroring the app's
SRC alpha and RMD-reduction alpha charts (tax-drag-jfp.md Figures 1 and 2). openpyxl writes the data,
formulas and charts; Excel (via pywin32) then recalculates and adds the titles, which it lays out properly."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter as CL

FIG = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(FIG, 'tax-drag-charts.xlsx')
F = 'Arial'
BLUE, BLACK, HDR = Font(name=F, color='0000FF'), Font(name=F), Font(name=F, bold=True)
NOTE = Font(name=F, size=9, italic=True)
USD, USD2 = '$#,##0;($#,##0);-', '$#,##0.00;($#,##0.00);-'
H = 16  # data header row

def load(name):
    with open(os.path.join(FIG, name)) as f:
        return list(csv.DictReader(f))

def build(wb, spec):
    ws = wb.create_sheet(spec['sheet'])
    ws['A1'] = spec['title']; ws['A1'].font = Font(name=F, bold=True, size=12)
    ws['A2'] = spec['note']; ws['A2'].font = NOTE
    ref = {}
    for i, (key, label, val, fmt, comment) in enumerate(spec['inputs']):
        r = 4 + i
        ws.cell(r, 1, label).font = BLACK
        c = ws.cell(r, 2, val); c.number_format = fmt
        c.font = BLACK if isinstance(val, str) and val.startswith('=') else BLUE
        if comment: ws.cell(r, 3, comment).font = NOTE
        ref[key] = f'$B${r}'

    cols = spec['cols']; idx = {k: i + 1 for i, (k, _, _) in enumerate(cols)}
    L = lambda k: CL(idx[k])
    for k, head, _ in cols:
        c = ws.cell(H, idx[k], head); c.font = HDR; c.border = Border(bottom=Side(style='thin'))
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='bottom')
    ws.row_dimensions[H].height = 45
    first = H + 1; last = first + len(spec['rows']) - 1
    for i, row in enumerate(spec['rows']):
        n = first + i
        for k, _, kind in cols:
            if kind == 'gap':   v, fmt, font = f'={L("roth_dist")}{n}-{L("taxdrag_dist")}{n}', USD2, BLACK
            elif kind == 'pv':  v, fmt, font = f'={L("gap")}{n}/(1+{ref["r"]})^({L("year")}{n}-{ref["base"]})', USD2, BLACK
            elif kind == 'cum': v, fmt, font = f'=SUM(${L("gap")}${first}:{L("gap")}{n})', USD, BLACK
            elif kind == 'cumpv': v, fmt, font = f'=SUM(${L("gap_pv")}${first}:{L("gap_pv")}{n})', USD, BLACK
            else:
                raw = row[k]
                v = None if raw == '' else (int(raw) if kind == 'int' else float(raw))
                fmt, font = {'int': '0', 'usd2': USD2}.get(kind, USD), BLUE
            c = ws.cell(n, idx[k], v); c.number_format = fmt; c.font = font
    # Helper for the chart's even-year vertical lines: 0 on even years, blank otherwise
    gc = len(cols) + 1
    c = ws.cell(H, gc, 'Gridline helper (even years)'); c.font = NOTE
    c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='bottom')
    for i, row in enumerate(spec['rows']):  # true blanks on odd years, so no line is drawn there
        if int(row['year']) % 2 == 0:
            ws.cell(first + i, gc, 0).font = Font(name=F, color='A6A6A6')
    tot = last + 2
    ws.cell(tot, 1, 'Total').font = HDR
    for k in ('gap', 'gap_pv'):
        c = ws[f'{L(k)}{tot}']; c.value = f'=SUM({L(k)}{first}:{L(k)}{last})'; c.number_format = USD; c.font = HDR
    ws.column_dimensions['A'].width = 24
    for j in range(2, len(cols) + 1):
        ws.column_dimensions[CL(j)].width = 14

    # Chart: balance lines on the left axis, cumulative-benefit bars on the right axis
    line = LineChart()
    for key, color, dash in spec['lines']:
        line.add_data(Reference(ws, min_col=idx[key], min_row=H, max_row=last), titles_from_data=True)
        s = line.series[-1]; s.smooth = False; s.marker.symbol = 'none'
        s.graphicalProperties.line.solidFill = color; s.graphicalProperties.line.width = 22000
        if dash: s.graphicalProperties.line.dashStyle = 'dash'
    line.set_categories(Reference(ws, min_col=idx['year'], min_row=first, max_row=last))
    line.y_axis.numFmt = '$#,##0'; line.y_axis.majorGridlines = None; line.x_axis.tickLblSkip = 2
    bar = BarChart()
    bar.add_data(Reference(ws, min_col=idx['cum'], min_row=H, max_row=last), titles_from_data=True)
    bar.series[0].graphicalProperties.solidFill = spec['bar'][0]
    bar.series[0].graphicalProperties.line.solidFill = spec['bar'][1]
    # The bars share the line axes here; Excel moves them to its secondary axis below. Writing the
    # secondary axis in openpyxl and then regrouping in Excel saves a file Excel cannot reopen.
    bar.gapWidth = 40
    for ax in (line.x_axis, line.y_axis):
        ax.delete = False
    line += bar
    line.legend.position = 'b'; line.legend.overlay = False
    line.width, line.height = 28, 15
    ws.add_chart(line, f'{CL(len(cols) + 2)}{H}')
    return dict(ws=spec['sheet'], tot=tot, fv=L('gap'), pv=L('gap_pv'), lines=spec['title_lines'],
                grid=f'{CL(gc)}{first}:{CL(gc)}{last}', years=f'{L("year")}{first}:{L("year")}{last}')

def common_cols(principal):
    return [('year', 'Year', 'int'), ('age', 'Age', 'int')] + principal + [
        ('roth_dist', 'Roth distribution', 'usd2'), ('taxdrag_dist', 'Tax-Drag distribution', 'usd2'),
        ('gap', 'Yearly gap', 'gap'), ('gap_pv', 'Yearly gap PV', 'pv'),
        ('cum', 'Tax-Drag Benefit - Cum', 'cum'), ('cum_pv', 'Tax-Drag Benefit - Cum PV', 'cumpv')]

SRC = dict(
    sheet='SRC alpha', rows=load('src-alpha-chart-data.csv'),
    title='SRC Alpha (Wrapper Migration) - Moves Taxable Cash into the Roth',
    note='Figure 1, tax-drag-jfp.md. K in a sheltered Roth (7%) vs. a taxable account (6.65%); RMDs 2039–2058, '
         'then a 10-year beneficiary annuity. Age is blank in beneficiary years.',
    inputs=[('k', 'Outside Tax Payment', 17966, USD, 'Conversion tax K'),
            ('r', 'Return', 0.07, '0.0%', 'Pre-tax return r; also the PV discount rate'),
            ('d', 'Tax Drag', 0.05, '0.0%', 'Share of return lost to tax in the taxable account'),
            ('toage', 'To Age', 94, '0', "Owner's age in the last distribution year (2058)"),
            ('life', 'Dist Years', 20, '0', 'Owner-life distribution years (2039–2058)'),
            ('m', 'Dist Method', 'RMD', '@', 'RMDs from age 75, then 10-year beneficiary annuity'),
            ('base', 'Base year', 2026, '0', 'PV is discounted to this year')],
    cols=common_cols([('roth_bal', 'Roth Balance', 'usd'), ('taxdrag_bal', 'Tax-Drag Balance', 'usd')]),
    lines=[('roth_bal', '1565C0', False), ('taxdrag_bal', 'D32F2F', False)],
    bar=('90CAF9', '1565C0'),
    title_lines=('SRC Alpha (Wrapper Migration)', ' - Moves Taxable Cash into the Roth',
                 'Outside Tax Payment: $17,966     Return: 7.0%     Tax Drag: 5.0%     To Age: 94     Dist Years: 20     Dist Method: RMD'))

RMD = dict(
    sheet='RMD alpha', rows=load('rmd-alpha-chart-data.csv'),
    title='RMD Reduction Alpha — Benefit of avoiding RMDs on $96,359 of after-distribution-tax principal',
    note='Figure 2, tax-drag-jfp.md. 75% of each avoided RMD is reinvested in parallel Roth (7%) and taxable (6.65%) '
         'balances 2039–2058; both then distribute as a 10-year beneficiary annuity. Age is blank in beneficiary years.',
    inputs=[('c', 'Conversion Amount', 121800, USD, 'C, the 22% bracket fill'),
            ('td', 'Displaced tax rate tD', 0.20887685, '0.0000%', 'Paper rounds to 20.89%; this value reproduces P = $96,359'),
            ('p', 'After-tax principal C(1−tD)', '=B4*(1-B5)', USD, None),
            ('r', 'Return', 0.07, '0.0%', 'Pre-tax return r; also the PV discount rate'),
            ('d', 'Tax Drag', 0.05, '0.0%', 'Share of return lost to tax in the taxable account'),
            ('rho', 'Reinvestment share', 0.75, '0%', 'Share of each avoided RMD reinvested'),
            ('toage', 'To Age', 94, '0', "Owner's age in the last distribution year (2058)"),
            ('life', 'Dist Years', 20, '0', 'Owner-life distribution years (2039–2058)'),
            ('m', 'Dist Method', 'RMD', '@', 'RMDs from age 75, then 10-year beneficiary annuity'),
            ('base', 'Base year', 2026, '0', 'PV is discounted to this year')],
    cols=common_cols([('after_tax_principal', 'After-Tax Principal C(1−tD)', 'usd'),
                      ('reinvest_roth_bal', 'Dist Reinvest - Roth', 'usd'),
                      ('reinvest_taxdrag_bal', 'Dist Reinvest - Tax-Drag', 'usd')]),
    lines=[('after_tax_principal', 'FF9800', True), ('reinvest_roth_bal', '1565C0', False),
           ('reinvest_taxdrag_bal', 'D32F2F', False)],
    bar=('81C784', '2E7D32'),
    title_lines=('RMD Reduction Alpha', ' — Benefit of avoiding RMDs on $96,359 of after-distribution-tax principal',
                 'Conversion Amount: $121,800     Return: 7.0%     Tax Drag: 5.0%     To Age: 94     Dist Years: 20     Dist Method: RMD'))

wb = Workbook(); wb.remove(wb.active)
charts = [build(wb, SRC), build(wb, RMD)]
wb.save(OUT)

# Excel: recalculate, then add the chart titles and the right-axis title
import win32com.client
xl = win32com.client.DispatchEx('Excel.Application'); xl.Visible = False; xl.DisplayAlerts = False
try:
    book = xl.Workbooks.Open(OUT); xl.CalculateFull()
    for c in charts:
        ws = book.Worksheets(c['ws'])
        fv, pv = ws.Range(f"{c['fv']}{c['tot']}").Text, ws.Range(f"{c['pv']}{c['tot']}").Text
        bold, rest, inputs = c['lines']
        line3 = f'Benefit FV: {fv}     PV: {pv}'
        text = f'{bold}{rest}\n{inputs}\n{line3}'
        ch = ws.ChartObjects(1).Chart
        ch.HasTitle = True; t = ch.ChartTitle; t.Text = text
        tr = t.Format.TextFrame2.TextRange
        tr.Font.Size = 10; tr.Font.Bold = 0; tr.Font.Name = F
        tr.GetCharacters(1, len(bold)).Font.Bold = -1   # pywin32 exposes Characters(start, len) as GetCharacters
        tr.GetCharacters(1, len(bold)).Font.Size = 12
        s3 = len(text) - len(line3) + 1
        tr.GetCharacters(s3, len('Benefit FV:')).Font.Bold = -1
        tr.GetCharacters(s3 + line3.index('PV:'), len('PV:')).Font.Bold = -1
        # Bars to the secondary (right) axis
        for i in range(1, ch.SeriesCollection().Count + 1):
            if ch.SeriesCollection(i).ChartType == 51:  # xlColumnClustered
                ch.SeriesCollection(i).AxisGroup = 2
        ax = ch.Axes(2, 2)  # xlValue, xlSecondary
        ax.TickLabels.NumberFormat = '$#,##0'; ax.HasMajorGridlines = False
        # Faint dashed vertical lines on even years, as in the app. Category gridlines would fall between
        # years (the bars need between-category placement), so draw them as error bars on a hidden series.
        ch.Axes(1, 1).TickLabelSpacing = 2
        ax1 = ch.Axes(2, 1); top = ax1.MaximumScale; ax1.MaximumScale = top
        g = ch.SeriesCollection().NewSeries()
        g.Values = ws.Range(c['grid']); g.XValues = ws.Range(c['years'])
        g.ChartType = 4; g.AxisGroup = 1  # xlLine, primary
        g.Format.Line.Visible = 0; g.MarkerStyle = -4142  # no line, no markers
        g.ErrorBar(1, 2, 1, top)  # xlY, xlPlusValues, xlFixedValue, full axis height
        eb = g.ErrorBars; eb.EndStyle = 2  # xlNoCap
        eb.Format.Line.ForeColor.RGB = 0xD9D9D9; eb.Format.Line.Weight = 0.75; eb.Format.Line.DashStyle = 4
        ch.Legend.LegendEntries(ch.Legend.LegendEntries().Count).Delete()
        ax.HasTitle = True; ax.AxisTitle.Text = 'Cumulative Benefit'
        ax.AxisTitle.Font.Bold = True; ax.AxisTitle.Font.Size = 9; ax.AxisTitle.Orientation = -4170  # reads top-down, as in the app
        print(c['ws'], '| FV', fv, '| PV', pv)
    book.Save(); book.Close(True)
finally:
    xl.Quit()
print(OUT)
